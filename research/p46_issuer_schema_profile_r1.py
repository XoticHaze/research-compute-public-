from __future__ import annotations
import io,json,re,traceback,urllib.request
from pathlib import Path
import openpyxl
from xml.etree import ElementTree as ET

URLS={
"SPY":"https://www.ssga.com/library-content/products/fund-data/etfs/us/navhist-us-en-spy.xlsx",
"TLT":"https://www.blackrock.com/varnish-api/blk-one01-product-data/product-data/api/v1/get-fund-document?appSubType=ISHARES&appType=PRODUCT_PAGE&component=fundDownload&locale=en_US&portfolioId=239454&targetSite=us-ishares&userType=individual",
"GLD":"https://api.spdrgoldshares.com/api/v1/historical-archive?exchange=NYSE&lang=en&product=gld",
}

def get(u):
 req=urllib.request.Request(u,headers={"User-Agent":"Mozilla/5.0 CC-Market-Research/1.0"})
 with urllib.request.urlopen(req,timeout=60) as r:return r.read()

def xlsx_profile(b):
 wb=openpyxl.load_workbook(io.BytesIO(b),read_only=True,data_only=True)
 out={}
 for ws in wb.worksheets:
  rows=[]
  max_row=ws.max_row if isinstance(ws.max_row,int) else 40
  for row in ws.iter_rows(min_row=1,max_row=min(max_row,60),values_only=True):
   rows.append([None if v is None else str(v) for v in row[:24]])
  out[ws.title]={"max_row":ws.max_row,"max_column":ws.max_column,"head":rows}
 return out

def xml_profile(b):
 root=ET.fromstring(b)
 ns={'ss':'urn:schemas-microsoft-com:office:spreadsheet'}
 out={}
 for ws in root.findall('.//ss:Worksheet',ns):
  name=ws.attrib.get('{urn:schemas-microsoft-com:office:spreadsheet}Name','sheet')
  all_rows=ws.findall('.//ss:Row',ns); rows=[]
  for row in all_rows[:60]:
   vals=[]
   for cell in row.findall('ss:Cell',ns)[:24]:
    d=cell.find('ss:Data',ns); vals.append(None if d is None else d.text)
   rows.append(vals)
  out[name]={"head":rows,"row_count":len(all_rows)}
 return out

def dates_from_text(obj):
 txt=json.dumps(obj,default=str)
 hits=sorted(set(re.findall(r'(?:19|20)\d{2}[-/]\d{1,2}[-/]\d{1,2}',txt)))
 return {"date_token_count":len(hits),"first_date_token":hits[0] if hits else None,"last_date_token":hits[-1] if hits else None}

def main():
 out={"schema":"research.p46_issuer_schema_profile_r1","parent":"P46","sources":{},"failures":{}}
 for ticker,u in URLS.items():
  try:
   b=get(u); kind='xlsx' if b[:2]==b'PK' else 'xml' if b.lstrip().startswith(b'<?xml') else 'other'
   prof=xlsx_profile(b) if kind=='xlsx' else xml_profile(b) if kind=='xml' else {"preview":b[:1000].decode('utf-8','ignore')}
   out['sources'][ticker]={"kind":kind,"bytes":len(b),"profile":prof,**dates_from_text(prof)}
  except Exception as e:
   out['failures'][ticker]={"error":repr(e),"traceback":traceback.format_exc()}
 Path('artifacts').mkdir(exist_ok=True)
 Path('artifacts/p46_issuer_schema_profile_r1.json').write_text(json.dumps(out,indent=2,sort_keys=True,default=str))
 print(json.dumps({"sources":{k:{"kind":v['kind'],"bytes":v['bytes'],"date_token_count":v['date_token_count'],"first_date_token":v['first_date_token'],"last_date_token":v['last_date_token'],"sheets":list(v['profile'])} for k,v in out['sources'].items()},"failures":out['failures']},sort_keys=True))
 if out['failures']: raise SystemExit(2)
if __name__=='__main__':main()
