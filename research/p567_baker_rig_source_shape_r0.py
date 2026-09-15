from __future__ import annotations

import hashlib
import io
import json
from pathlib import Path

import requests
from openpyxl import load_workbook

SOURCES = {
    "current_2026_09_04": "https://rigcount.bakerhughes.com/static-files/2da8181e-4b1c-4f75-9ad8-44854f4fc106",
    "archive_2013_2025_08": "https://rigcount.bakerhughes.com/static-files/e98bcf83-c458-4a88-8f35-4ac4d77628bb",
}
UA = {"User-Agent": "XoticHaze Research xotichaze@users.noreply.github.com"}
TERMS = ("oil", "gas", "u.s.", "united states", "publish", "date", "drill for")


def scalar(v):
    if v is None:
        return None
    if hasattr(v, "isoformat"):
        try:
            return v.isoformat()
        except Exception:
            pass
    return str(v)


def inspect_book(content: bytes) -> dict:
    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    out = {"sheets": []}
    for ws in wb.worksheets:
        sheet = {
            "title": ws.title,
            "max_row": ws.max_row,
            "max_column": ws.max_column,
            "head_rows": [],
            "term_rows": [],
        }
        for ri, row in enumerate(ws.iter_rows(values_only=True), start=1):
            vals = [scalar(v) for v in row]
            compact = [v for v in vals if v not in (None, "")]
            if ri <= 12 and compact:
                sheet["head_rows"].append({"row": ri, "values": compact[:30]})
            joined = " | ".join(compact).lower()
            if joined and any(t in joined for t in TERMS) and len(sheet["term_rows"]) < 30:
                sheet["term_rows"].append({"row": ri, "values": compact[:40]})
            if ri >= 300 and len(sheet["term_rows"]) >= 30:
                break
        out["sheets"].append(sheet)
    return out


def main() -> None:
    session = requests.Session()
    result = {
        "schema": "research.p567_baker_rig_source_shape_r0",
        "parent": "P07",
        "child": "P567",
        "decision": "SOURCE_SHAPE_ONLY",
        "sources": {},
        "boundaries": {
            "alpha_claim": False,
            "portfolio_ranking": False,
            "allocation_authority": False,
            "runtime": False,
            "broker": False,
            "live_trading": False,
        },
    }
    for name, url in SOURCES.items():
        item = {"url": url}
        try:
            r = session.get(url, headers=UA, timeout=(20, 120), allow_redirects=True)
            item.update({
                "status": r.status_code,
                "final_url": r.url,
                "bytes": len(r.content),
                "content_type": r.headers.get("content-type"),
                "sha256": hashlib.sha256(r.content).hexdigest(),
            })
            r.raise_for_status()
            item["workbook"] = inspect_book(r.content)
        except Exception as exc:
            item["exception"] = repr(exc)
        result["sources"][name] = item

    Path("artifacts").mkdir(exist_ok=True)
    out = Path("artifacts/p567_baker_rig_source_shape_r0.json")
    out.write_text(json.dumps(result, indent=2, sort_keys=True))
    print(json.dumps(result, sort_keys=True))


if __name__ == "__main__":
    main()
