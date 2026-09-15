from __future__ import annotations

import io
import json
import math
from pathlib import Path

import pandas as pd
import requests
import yfinance as yf
from openpyxl import load_workbook

ARCHIVE_URL = "https://rigcount.bakerhughes.com/static-files/e98bcf83-c458-4a88-8f35-4ac4d77628bb"
CURRENT_URL = "https://rigcount.bakerhughes.com/static-files/2da8181e-4b1c-4f75-9ad8-44854f4fc106"
UA = {"User-Agent": "XoticHaze Research xotichaze@users.noreply.github.com"}
ASSETS = ["OIH", "XLE", "XOP", "SPY"]
COST = 25.0 / 10000.0
FOLDS = [
    ("2014-01-01", "2016-12-31"),
    ("2017-01-01", "2019-12-31"),
    ("2020-01-01", "2022-12-31"),
    ("2023-01-01", "2026-09-11"),
]


def load_us_oil(url: str) -> pd.Series:
    r = requests.get(url, headers=UA, timeout=(20, 120))
    r.raise_for_status()
    wb = load_workbook(io.BytesIO(r.content), read_only=True, data_only=True)
    ws = wb["NAM Weekly"]
    header = None
    data = []
    for row in ws.iter_rows(values_only=True):
        vals = list(row)
        if header is None:
            norm = [str(v).strip() if v is not None else "" for v in vals]
            if "US_PublishDate" in norm and "Rig Count Value" in norm:
                header = norm
            continue
        rec = dict(zip(header, vals))
        if str(rec.get("Country") or "").strip().upper() != "UNITED STATES":
            continue
        if str(rec.get("DrillFor") or "").strip().upper() != "OIL":
            continue
        dt = pd.to_datetime(rec.get("US_PublishDate"), errors="coerce")
        val = pd.to_numeric(rec.get("Rig Count Value"), errors="coerce")
        if pd.notna(dt) and pd.notna(val):
            data.append((pd.Timestamp(dt).normalize(), float(val)))
    if not data:
        raise RuntimeError(f"no U.S. oil rig rows from {url}")
    df = pd.DataFrame(data, columns=["date", "value"])
    return df.groupby("date", sort=True).value.sum().sort_index()


def cagr(r: pd.Series) -> float | None:
    r = r.dropna()
    if len(r) < 2:
        return None
    years = (r.index[-1] - r.index[0]).days / 365.25
    total = float((1 + r).prod())
    return None if years <= 0 or total <= 0 else total ** (1 / years) - 1


def mdd(r: pd.Series) -> float | None:
    r = r.dropna()
    if r.empty:
        return None
    eq = (1 + r).cumprod()
    return float((eq / eq.cummax() - 1).min())


def stats(r: pd.Series) -> dict:
    r = r.dropna()
    return {
        "cagr": cagr(r),
        "max_drawdown": mdd(r),
        "vol": None if len(r) < 2 else float(r.std() * math.sqrt(252)),
        "days": int(len(r)),
    }


def evaluate(d: pd.DataFrame, start: str, end: str, matched_weight: float) -> dict:
    z = d.loc[start:end].copy()
    matched = matched_weight * z.OIH + (1 - matched_weight) * z.XLE
    sc = cagr(z.strategy)
    mc = cagr(matched)
    xc = cagr(z.XLE)
    return {
        "window": [start, end],
        "strategy": stats(z.strategy),
        "matched_static": stats(matched),
        "OIH": stats(z.OIH),
        "XLE": stats(z.XLE),
        "XOP": stats(z.XOP),
        "SPY": stats(z.SPY),
        "matched_excess_cagr": None if sc is None or mc is None else sc - mc,
        "xle_excess_cagr": None if sc is None or xc is None else sc - xc,
        "xop_excess_cagr": None if sc is None else sc - cagr(z.XOP),
        "spy_excess_cagr": None if sc is None else sc - cagr(z.SPY),
        "oih_excess_cagr": None if sc is None else sc - cagr(z.OIH),
        "mean_on_exposure": None if z.empty else float(z.on.mean()),
        "switches": int(z.switch.sum()) if not z.empty else 0,
    }


def main() -> None:
    archive = load_us_oil(ARCHIVE_URL)
    current = load_us_oil(CURRENT_URL)
    overlap = pd.concat([archive.rename("archive"), current.rename("current")], axis=1, join="inner").dropna()
    mismatch = (overlap.archive - overlap.current).abs() > 1e-9
    mismatch_rate = None if overlap.empty else float(mismatch.mean())
    max_abs_diff = None if overlap.empty else float((overlap.archive - overlap.current).abs().max())

    stitched = archive.copy()
    stitched = pd.concat([stitched, current[current.index > stitched.index.max()]])
    stitched = stitched[~stitched.index.duplicated(keep="first")].sort_index()
    signal = (stitched > stitched.shift(13)).astype(float).iloc[13:]

    raw = yf.download(ASSETS, start="2013-01-01", end="2026-09-12", auto_adjust=True, progress=False, group_by="column", threads=False)
    close = raw["Close"] if isinstance(raw.columns, pd.MultiIndex) else raw
    rets = close[ASSETS].dropna(how="any").pct_change(fill_method=None).dropna(how="any")

    on = pd.Series(index=rets.index, dtype=float)
    events = []
    for pub_date, state in signal.items():
        eligible = rets.index[rets.index > pub_date]
        if len(eligible) == 0:
            continue
        effective = eligible[0]
        on.loc[effective] = float(state)
        events.append({"publish_date": str(pub_date.date()), "effective_date": str(effective.date()), "oil_rigs": float(stitched.loc[pub_date]), "oil_rigs_13w_prior": float(stitched.shift(13).loc[pub_date]), "on": float(state)})
    on = on.ffill()
    first = on.first_valid_index()
    if first is None:
        raise RuntimeError("no causal P567 signal became effective")

    d = rets.loc[first:].copy()
    d["on"] = on.loc[d.index]
    d = d.dropna(subset=["on"])
    d["switch"] = d.on.diff().abs().fillna(0)
    d["strategy"] = d.on * d.OIH + (1 - d.on) * d.XLE - d.switch * COST
    matched_weight = float(d.on.mean())
    overall = evaluate(d, str(d.index.min().date()), "2026-09-11", matched_weight)
    folds = [evaluate(d, a, b, matched_weight) for a, b in FOLDS]
    positive_folds = sum(1 for f in folds if f["matched_excess_cagr"] is not None and f["matched_excess_cagr"] > 0)

    source_safe = mismatch_rate is not None and mismatch_rate <= 0.01
    supported = bool(
        source_safe
        and overall["matched_excess_cagr"] is not None
        and overall["matched_excess_cagr"] > 0
        and overall["xle_excess_cagr"] is not None
        and overall["xle_excess_cagr"] > 0
        and positive_folds >= 3
    )
    if not source_safe:
        decision = "P567_SOURCE_REVISION_RISK_NO_ALPHA_CONCLUSION"
    else:
        decision = "P567_SUPPORTED" if supported else "P567_NOT_SUPPORTED_NO_RESCUE"

    out = {
        "schema": "research.p567_baker_rig_oil_services_r1",
        "parent": "P07",
        "child": "P567",
        "claim": "Causally published U.S. oil drilling activity acceleration can identify periods when oilfield services outperform broad energy after switching costs.",
        "frozen_contract": {
            "signal": "current U.S. oil rig count > count 13 published observations earlier",
            "on": "OIH",
            "off": "XLE",
            "availability": "first U.S. trading session strictly after Baker Hughes US_PublishDate",
            "full_switch_cost_bps": 25.0,
            "matched_control": "static OIH/XLE at full-sample mean OIH exposure",
            "folds": FOLDS,
            "pass_gate": "full matched excess >0, full XLE excess >0, >=3/4 positive matched folds, overlap mismatch rate <=1%",
            "no_parameter_rescue": True,
        },
        "source": {
            "archive_first": str(archive.index.min().date()),
            "archive_last": str(archive.index.max().date()),
            "archive_weeks": int(len(archive)),
            "current_first": str(current.index.min().date()),
            "current_last": str(current.index.max().date()),
            "current_weeks": int(len(current)),
            "stitched_first": str(stitched.index.min().date()),
            "stitched_last": str(stitched.index.max().date()),
            "stitched_weeks": int(len(stitched)),
            "overlap_weeks": int(len(overlap)),
            "overlap_mismatches": int(mismatch.sum()) if len(overlap) else None,
            "overlap_mismatch_rate": mismatch_rate,
            "overlap_max_abs_count_difference": max_abs_diff,
        },
        "signal_diagnostics": {"effective_events": len(events), "mean_on_exposure": matched_weight, "switches": int(d.switch.sum())},
        "overall": overall,
        "folds": folds,
        "positive_fold_count": positive_folds,
        "decision": decision,
        "boundaries": {"portfolio_ranking": False, "allocation_authority": False, "runtime": False, "broker": False, "live_trading": False},
    }
    Path("artifacts").mkdir(exist_ok=True)
    Path("artifacts/p567_baker_rig_oil_services_r1.json").write_text(json.dumps(out, indent=2, sort_keys=True))
    print(json.dumps(out, sort_keys=True))


if __name__ == "__main__":
    main()
