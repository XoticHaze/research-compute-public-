from __future__ import annotations

import io
import json
from pathlib import Path

import pandas as pd
import requests
from openpyxl import load_workbook

ARCHIVE_URL = "https://rigcount.bakerhughes.com/static-files/e98bcf83-c458-4a88-8f35-4ac4d77628bb"
CURRENT_URL = "https://rigcount.bakerhughes.com/static-files/2da8181e-4b1c-4f75-9ad8-44854f4fc106"
UA = {"User-Agent": "XoticHaze Research xotichaze@users.noreply.github.com"}


def load_us_oil(url: str) -> pd.Series:
    r = requests.get(url, headers=UA, timeout=(20, 120))
    r.raise_for_status()
    wb = load_workbook(io.BytesIO(r.content), read_only=True, data_only=True)
    ws = wb["NAM Weekly"]
    rows = ws.iter_rows(values_only=True)
    header = None
    data = []
    for row in rows:
        vals = list(row)
        if header is None:
            norm = [str(v).strip() if v is not None else "" for v in vals]
            if "US_PublishDate" in norm and "Rig Count Value" in norm:
                header = norm
            continue
        rec = dict(zip(header, vals))
        if str(rec.get("Country") or "").strip().upper() != "UNITED STATES":
            continue
        if str(rec.get("DrillFor") or "").strip().upper() != "OIL":
            continue
        dt = pd.to_datetime(rec.get("US_PublishDate"), errors="coerce")
        value = pd.to_numeric(rec.get("Rig Count Value"), errors="coerce")
        if pd.notna(dt) and pd.notna(value):
            data.append((pd.Timestamp(dt).normalize(), float(value)))
    if not data:
        raise RuntimeError(f"no U.S. Oil rows parsed from {url}")
    df = pd.DataFrame(data, columns=["date", "value"])
    return df.groupby("date", sort=True).value.sum().sort_index()


def main() -> None:
    archive = load_us_oil(ARCHIVE_URL)
    current = load_us_oil(CURRENT_URL)
    overlap = pd.concat([archive.rename("archive"), current.rename("current")], axis=1, join="inner").dropna()
    overlap["diff"] = overlap.current - overlap.archive
    overlap["abs_diff"] = overlap["diff"].abs()
    mismatch = overlap[overlap.abs_diff > 1e-9]
    mismatch_rate = None if len(overlap) == 0 else len(mismatch) / len(overlap)
    out = {
        "schema": "research.p567_baker_rig_overlap_audit_r0",
        "parent": "P07",
        "child": "P567",
        "archive": {"first": str(archive.index.min().date()), "last": str(archive.index.max().date()), "weeks": int(len(archive))},
        "current": {"first": str(current.index.min().date()), "last": str(current.index.max().date()), "weeks": int(len(current))},
        "overlap": {
            "first": None if overlap.empty else str(overlap.index.min().date()),
            "last": None if overlap.empty else str(overlap.index.max().date()),
            "weeks": int(len(overlap)),
            "mismatches": int(len(mismatch)),
            "mismatch_rate": mismatch_rate,
            "max_abs_count_difference": None if overlap.empty else float(overlap.abs_diff.max()),
            "mismatch_samples": [
                {"date": str(idx.date()), "archive": float(row.archive), "current": float(row.current), "diff": float(row["diff"])}
                for idx, row in mismatch.head(20).iterrows()
            ],
        },
        "decision": "OVERLAP_STABLE" if mismatch_rate is not None and mismatch_rate <= 0.01 else "OVERLAP_REVISION_RISK",
        "boundaries": {"alpha_claim": False, "portfolio_ranking": False, "allocation_authority": False, "runtime": False, "broker": False, "live_trading": False},
    }
    Path("artifacts").mkdir(exist_ok=True)
    Path("artifacts/p567_baker_rig_overlap_audit_r0.json").write_text(json.dumps(out, indent=2, sort_keys=True))
    print(json.dumps(out, sort_keys=True))


if __name__ == "__main__":
    main()
