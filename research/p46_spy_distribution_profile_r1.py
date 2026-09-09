import io,json,hashlib,urllib.request
from pathlib import Path
import openpyxl
U='https://www.ssga.com/library-content/products/fund-data/etfs/us/spdr-etf-historical-distributions.xlsx'
def main():
 b=urllib.request.urlopen(urllib.request.Request(U,headers={'User-Agent':'Mozilla/5.0 CC-Market-Research/1.0'}),timeout=60).read();w=openpyxl.load_workbook(io.BytesIO(b),read_only=True,data_only=True);o={'schema':'research.p46_spy_distribution_profile_r1','parent':'P46','url':U,'bytes':len(b),'sha256':hashlib.sha256(b).hexdigest(),'sheets':{}}
 for s in w.sheetnames:
  ws=w[s];rows=[]
  for r in ws.iter_rows(min_row=1,max_row=min(ws.max_row or 80,80),values_only=True):rows.append([None if x is None else str(x) for x in r[:24]])
  o['sheets'][s]={'rows':ws.max_row,'columns':ws.max_column,'head':rows}
 Path('artifacts').mkdir(exist_ok=True);Path('artifacts/p46_spy_distribution_profile_r1.json').write_text(json.dumps(o,indent=2,sort_keys=True));print(json.dumps({'bytes':len(b),'sha256':o['sha256'],'sheets':{k:{'rows':v['rows'],'columns':v['columns']} for k,v in o['sheets'].items()}},sort_keys=True))
if __name__=='__main__':main()
