from __future__ import annotations
import io,json,urllib.request
from pathlib import Path
from datetime import date
import openpyxl,pandas as pd
NAV='https://www.ssga.com/library-content/products/fund-data/etfs/us/navhist-us-en-spy.xlsx'
DIST='https://www.ssga.com/library-content/products/fund-data/etfs/us/spdr-etf-historical-distributions.xlsx'
TARGET=21.45
TOL=0.10

def get(u):
 return urllib.request.urlopen(urllib.request.Request(u,headers={'User-Agent':'Mozilla/5.0 CC-Market-Research/1.0'}),timeout=60).read()
def navs(b):
 ws=openpyxl.load_workbook(io.BytesIO(b),read_only=True,data_only=True)['navhist']; rows=[list(r) for r in ws.iter_rows(values_only=True)]; i=next(i for i,r in enumerate(rows) if r and r[0]=='Date'); h=rows[i]; di=h.index('Date'); ni=h.index('NAV'); x=pd.DataFrame({'d':[r[di] for r in rows[i+1:]],'n':[r[ni] for r in rows[i+1:]]}); x.d=pd.to_datetime(x.d,errors='coerce');x.n=pd.to_numeric(x.n,errors='coerce');x=x.dropna().drop_duplicates('d').sort_values('d');return dict(zip(x.d.dt.date,x.n.astype(float)))
def dists(b):
 ws=openpyxl.load_workbook(io.BytesIO(b),read_only=True,data_only=True)['dividend']; rows=[list(r) for r in ws.iter_rows(values_only=True)]; h=rows[0]; ti=h.index('TICKER'); ei=h.index('EX-DATE'); cols=[h.index('DIVIDEND ($)'),h.index('SHORT TERM CAPITAL GAIN ($)'),h.index('LONG TERM CAPITAL GAIN ($)')]; out=[]
 for r in rows[1:]:
  if str(r[ti]).strip()!='SPY':continue
  d=pd.to_datetime(r[ei],errors='coerce');
  if pd.isna(d):continue
  amt=sum(float(r[j]) if r[j] not in (None,'') else 0.0 for j in cols);out.append((d.date(),amt))
 return sorted(out)
def on(nav,d):
 k=max(x for x in nav if x<=d);return k,nav[k]
def main():
 n=navs(get(NAV)); ds=dists(get(DIST)); start=date(2022,9,30); end=date(2023,9,30); sd,s=on(n,start); ed,e=on(n,end); sh=1.0; used=[]
 for d,a in ds:
  if start<d<=end:
   nd,nv=on(n,d); sh*=1+a/nv;used.append({'ex_date':str(d),'nav_date':str(nd),'amount_per_share':a,'nav':nv})
 ret=100*(sh*e/s-1);err=abs(ret-TARGET);dec='P46_SPY_NAV_IDENTITY_VALIDATED' if err<=TOL else 'P46_SPY_NAV_IDENTITY_NOT_VALIDATED';o={'schema':'research.p46_spy_identity_validation_r1','parent':'P46','contract':{'period_end':'2023-09-30','one_year_sec_nav_total_return_pct':TARGET,'reinvestment':'all ordinary/ST/LT distributions at issuer NAV on ex-date','predeclared_tolerance_pp':TOL,'no_model_parameter_or_cost_changes':True},'nav_rows':len(n),'distribution_rows_total':len(ds),'start_date_used':str(sd),'end_date_used':str(ed),'start_nav':s,'end_nav':e,'distributions_used':used,'reconstructed_nav_total_return_pct':ret,'abs_error_pp':err,'decision':dec};Path('artifacts').mkdir(exist_ok=True);Path('artifacts/p46_spy_identity_validation_r1.json').write_text(json.dumps(o,indent=2,sort_keys=True));print(json.dumps(o,sort_keys=True))
if __name__=='__main__':main()
